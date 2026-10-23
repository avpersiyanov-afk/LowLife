# -*- coding: utf-8 -*-
"""
Сохранение списка задач «Задач из почты» в .xlsx через openpyxl — для
отдельного приложения (apps/mail), где нет .NET и Excel может быть не
нужен для записи. Таблица та же, что у email_tasks_excel (COM Excel, в
Revit): колонки и статусы — email_tasks_core, скрытый лист «Списки» со
статусами и «сегодня», выпадающий список статусов, условное
форматирование (готово — серым, просрочено — красным, высокий приоритет —
жёлтым), автофильтр, закреплённая шапка.

Формулы в файле .xlsx всегда хранятся по-английски, и Excel сам переводит
их в язык интерфейса, — поэтому здесь не нужны обходы локали, как у COM.
"""

import datetime
import os

from lowlife.email_tasks_core import (
    COLUMNS, STATUSES, STATUS_DONE, OUTPUT_FOLDER_NAME, PRIORITY_HIGH, output_file_name, table_rows
)

DATE_FORMAT = u"dd.mm.yyyy"
DATETIME_FORMAT = u"dd.mm.yyyy hh:mm"
EXTRA_STATUS_ROWS = 200  # выпадающий список статусов — с запасом под ручные задачи


class ExcelError(Exception):
    def __init__(self, message):
        Exception.__init__(self, message)
        self.message = message


def documents_folder():
    """Папка «Документы» (на Windows — с учётом переноса в OneDrive)."""
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes
            buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
            # CSIDL_PERSONAL = 5, SHGFP_TYPE_CURRENT = 0
            if ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf) == 0 and buf.value:
                return buf.value
        except Exception:
            pass
    return os.path.join(os.path.expanduser(u"~"), u"Documents")


def default_output_path(now=None):
    folder = os.path.join(documents_folder(), OUTPUT_FOLDER_NAME)
    if not os.path.isdir(folder):
        os.makedirs(folder)
    return os.path.join(folder, output_file_name(now or datetime.datetime.now()))


def _col(idx):
    return u"ABCDEFGHIJKLMNOPQRSTUVWXYZ"[idx]


def _define_name(wb, name, ref):
    from openpyxl.workbook.defined_name import DefinedName
    defined = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names[name] = defined  # openpyxl 3.1+
    except TypeError:
        wb.defined_names.append(defined)  # openpyxl 3.0


def build_workbook(rows):
    """Книга openpyxl с таблицей задач (rows — см. email_tasks_core.normalize_task)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = u"Задачи"

    # --- скрытый лист со списком статусов и «сегодня»
    lists = wb.create_sheet(u"Списки")
    for i, status in enumerate(STATUSES, start=1):
        lists.cell(row=i, column=1, value=status)
    lists[u"B1"] = u"=TODAY()"
    lists.sheet_state = "hidden"
    _define_name(wb, u"StatusList", u"'Списки'!$A$1:$A${}".format(len(STATUSES)))
    _define_name(wb, u"TodayDate", u"'Списки'!$B$1")

    # --- данные
    ws.append([title for title, _width in COLUMNS])
    for values in table_rows(rows):
        ws.append(values)

    last_col = _col(len(COLUMNS) - 1)
    last_row = len(rows) + 1

    font = Font(name=u"Arial", size=10)
    header_font = Font(name=u"Arial", size=10, bold=True)
    header_fill = PatternFill(u"solid", fgColor=u"DDEBF7")
    side = Side(style=u"thin", color=u"BFBFBF")
    border = Border(left=side, right=side, top=side, bottom=side)

    for c, (_title, width) in enumerate(COLUMNS):
        ws.column_dimensions[_col(c)].width = width

    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=last_row), start=1):
        for c, cell in enumerate(row):
            cell.border = border
            if r == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical=u"center", wrap_text=True)
                continue
            cell.font = font
            cell.alignment = Alignment(vertical=u"top", wrap_text=True,
                                       horizontal=u"center" if c == 0 else None)
            if c == 3 and cell.value is not None:
                cell.number_format = DATE_FORMAT
            elif c == 9 and cell.value is not None:
                cell.number_format = DATETIME_FORMAT

    # --- выпадающий список статусов
    validation = DataValidation(type=u"list", formula1=u"=StatusList", allow_blank=True)
    validation.add(u"F2:F{}".format(last_row + EXTRA_STATUS_ROWS))
    ws.add_data_validation(validation)

    # --- условное форматирование; порядок — приоритет: «Готово» со StopIfTrue
    # гасит красный/жёлтый у завершённых задач
    if rows:
        ws.conditional_formatting.add(
            u"A2:{}{}".format(last_col, last_row),
            FormulaRule(formula=[u'$F2="{}"'.format(STATUS_DONE)], stopIfTrue=True,
                        font=Font(color=u"808080"),
                        fill=PatternFill(u"solid", bgColor=u"F2F2F2")))
        ws.conditional_formatting.add(
            u"D2:D{}".format(last_row),
            FormulaRule(formula=[u'AND($D2<>"",$D2<TodayDate,$F2<>"{}")'.format(STATUS_DONE)],
                        font=Font(color=u"C00000", bold=True),
                        fill=PatternFill(u"solid", bgColor=u"FFC7CE")))
        ws.conditional_formatting.add(
            u"E2:E{}".format(last_row),
            FormulaRule(formula=[u'$E2="{}"'.format(PRIORITY_HIGH)],
                        fill=PatternFill(u"solid", bgColor=u"FFEB9C")))

    ws.auto_filter.ref = u"A1:{}{}".format(last_col, last_row)
    ws.freeze_panes = u"A2"
    try:
        wb.calculation.fullCalcOnLoad = True  # пересчитать =TODAY() при открытии
    except Exception:
        pass
    return wb


def save_tasks(rows, path, open_file=True):
    """Сохраняет книгу в path и (open_file) открывает её программой для
    .xlsx по умолчанию. Возвращает список некритичных предупреждений."""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise ExcelError(u"Не найден модуль openpyxl — без него Excel-файл не записать.")
    try:
        build_workbook(rows).save(path)
    except Exception as ex:
        raise ExcelError(u"Не удалось сохранить Excel:\n{}\n\n{}".format(path, ex))

    warnings = []
    if open_file:
        try:
            open_path(path)
        except Exception as ex:
            warnings.append(u"Файл сохранён, но не открылся: {}".format(ex))
    return warnings


def open_path(path):
    """Открыть файл или папку программой по умолчанию."""
    if os.name == "nt":
        os.startfile(path)  # noqa — есть только на Windows
    else:
        import subprocess
        subprocess.Popen([u"xdg-open", path])
